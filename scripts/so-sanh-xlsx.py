"""So hai tep xlsx tep loi XML3176: tung sheet, tung o (gia tri + kieu so/chu), numFmt cua o so,
tieu de (dam/can giua), do rong cot. In khac biet; ma thoat 1 neu co khac biet NGOAI o bat dau '='.

Chay: python scripts/so-sanh-xlsx.py <cu.xlsx> <moi.xlsx>
"""
import itertools
import re
import sys
import zipfile

import openpyxl

sys.stdout.reconfigure(encoding='utf-8')


def do_rong(tep):
    """{ten sheet: {so cot: do rong}} doc thang <cols> (read_only cua openpyxl khong co)."""
    z = zipfile.ZipFile(tep)
    wb = z.read('xl/workbook.xml').decode('utf-8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    dich = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
    dich.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    ra = {}
    for ten, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        duong = dich[rid].lstrip('/')
        duong = duong if duong.startswith('xl/') else 'xl/' + duong
        dau = z.open(duong).read(200000).decode('utf-8', 'ignore')
        ra[ten] = {}
        for el in re.findall(r'<col[^>]*/?>', dau):
            # chi so cot co customWidth (do rong co chu y); cot chi mang style la do rong mac dinh cua Excel
            if not re.search(r'customWidth="(1|true)"', el):
                continue
            ra[ten][int(re.search(r'min="(\d+)"', el).group(1))] = float(re.search(r'width="([\d.]+)"', el).group(1))
    return ra


def kieu(o):
    return 'trong' if o.value in (None, '') else ('so' if o.data_type == 'n' else 'chu')


def main(cu, moi):
    a = openpyxl.load_workbook(cu, read_only=True)
    b = openpyxl.load_workbook(moi, read_only=True)
    loi, cong_thuc = [], []
    if a.sheetnames != b.sheetnames:
        loi.append(f'Ten/thu tu sheet: {a.sheetnames} != {b.sheetnames}')
    ra_a, ra_b = do_rong(cu), do_rong(moi)
    for ten in a.sheetnames:
        if ten not in b.sheetnames:
            continue
        if ra_a.get(ten) != ra_b.get(ten):
            loi.append(f'[{ten}] do rong: {ra_a.get(ten)} != {ra_b.get(ten)}')
        da, db = a[ten].iter_rows(), b[ten].iter_rows()
        so_dong_a = so_dong_b = 0
        for r, (ha, hb) in enumerate(itertools.zip_longest(da, db, fillvalue=()), start=1):
            so_dong_a += 1 if ha else 0
            so_dong_b += 1 if hb else 0
            n = max(len(ha), len(hb))
            ha = list(ha) + [None] * (n - len(ha))
            hb = list(hb) + [None] * (n - len(hb))
            for c, (x, y) in enumerate(zip(ha, hb), start=1):
                vx = None if x is None or x.value == '' else x.value
                vy = None if y is None or y.value == '' else y.value
                if x is not None and x.data_type == 'f':
                    cong_thuc.append(f'[{ten}] R{r}C{c}: cu la cong thuc {vx!r}, moi {vy!r}')
                    continue
                if vx != vy or (x is not None and y is not None and kieu(x) != kieu(y)):
                    loi.append(f'[{ten}] R{r}C{c}: {vx!r} ({kieu(x) if x else "-"}) != {vy!r} ({kieu(y) if y else "-"})')
                elif x is not None and y is not None and kieu(x) == 'so' and x.number_format != y.number_format:
                    loi.append(f'[{ten}] R{r}C{c}: numFmt {x.number_format!r} != {y.number_format!r}')
                elif r == 1 and x is not None and y is not None and (
                        bool(x.font.b) != bool(y.font.b) or x.alignment.horizontal != y.alignment.horizontal):
                    loi.append(f'[{ten}] tieu de C{c}: dam/can {x.font.b}/{x.alignment.horizontal} != {y.font.b}/{y.alignment.horizontal}')
                if len(loi) > 50:
                    break
        if so_dong_a != so_dong_b:
            loi.append(f'[{ten}] so dong lech: cu {so_dong_a}, moi {so_dong_b}')
        print(f'{ten:16} {max(so_dong_a, so_dong_b) - 1:>8} dong du lieu')
    print(f'\nO cong thuc o ban cu (khac biet CO Y): {len(cong_thuc)}')
    for d in cong_thuc[:20]:
        print('  ' + d)
    print(f'Khac biet KHAC: {len(loi)}')
    for d in loi[:50]:
        print('  ' + d)
    return 1 if loi else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], sys.argv[2]))
